import io,json,csv
from pathlib import Path
from bs4 import BeautifulSoup
from pypdf import PdfReader
from docx import Document
from docx.table import Table
from pptx import Presentation
from openpyxl import load_workbook
from striprtf.striprtf import rtf_to_text

SUPPORTED={'.txt','.md','.csv','.json','.html','.htm','.pdf','.docx','.pptx','.xlsx','.rtf'}
def extract(name, data):
    ext=Path(name).suffix.lower()
    stream=io.BytesIO(data)
    if ext=='.pdf':
        reader=PdfReader(stream)
        if len(reader.pages)>500: raise ValueError('PDF exceeds 500 pages.')
        return '\n'.join(p.extract_text() or '' for p in reader.pages)
    if ext=='.docx':
        d=Document(stream)
        return '\n'.join('\n'.join(' | '.join(c.text for c in row.cells) for row in block.rows) if isinstance(block,Table) else block.text for block in d.iter_inner_content())
    if ext=='.pptx': return '\n'.join(s.text for slide in Presentation(stream).slides for s in slide.shapes if s.has_text_frame)
    if ext=='.xlsx':
        wb=load_workbook(stream,read_only=True,data_only=True)
        try: return '\n'.join(f'{sheet.title}: '+' | '.join(str(c) for c in row if c is not None) for sheet in wb for row in sheet.iter_rows(values_only=True))
        finally: wb.close()
    text=data.decode('utf-8-sig')
    if ext in {'.html','.htm'}:
        soup=BeautifulSoup(text,'html.parser')
        for s in soup(['script','style']): s.decompose()
        return soup.get_text(' ',strip=True)
    if ext=='.json': return json.dumps(json.loads(text),ensure_ascii=False,indent=2)
    if ext=='.csv': return '\n'.join(' | '.join(r) for r in csv.reader(io.StringIO(text)))
    if ext=='.rtf': return rtf_to_text(text)
    return text
