import io,json,os,importlib
from pathlib import Path
import pytest
from fastapi.testclient import TestClient
from docx import Document
from pptx import Presentation
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
from app.engine import chunks,retrieve

@pytest.fixture
def client(tmp_path,monkeypatch):
    monkeypatch.setenv('RAG_DATA_DIR',str(tmp_path))
    monkeypatch.delenv('VERCEL',raising=False)
    monkeypatch.delenv('DATABASE_URL',raising=False)
    if os.environ.get('TEST_DATABASE_URL'):
        import psycopg
        from app.storage import initialize
        url=os.environ['TEST_DATABASE_URL']
        with psycopg.connect(url) as connection:
            assert connection.info.dbname.endswith('_test'), 'Use a disposable database with a name ending in _test'
            initialize(connection)
            connection.execute('DELETE FROM chunks')
            connection.execute('DELETE FROM documents')
        monkeypatch.setenv('DATABASE_URL',url)
    import app.main as main
    importlib.reload(main)
    with TestClient(main.app) as c: yield c

def payload(ext):
    text='Employees receive twenty five vacation days annually. Submit leave requests to your manager.'
    out=io.BytesIO()
    if ext=='.docx':
        d=Document();d.add_paragraph(text);d.save(out)
    elif ext=='.pptx':
        d=Presentation();s=d.slides.add_slide(d.slide_layouts[1]);s.shapes.title.text=text;d.save(out)
    elif ext=='.xlsx':
        w=Workbook();w.active.append([text]);w.save(out)
    elif ext=='.pdf':
        w=PdfWriter();page=w.add_blank_page(width=612,height=792)
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):w._add_object(font)})})
        stream=DecodedStreamObject();stream.set_data(f'BT /F1 12 Tf 20 700 Td ({text}) Tj ET'.encode());page[NameObject('/Contents')]=w._add_object(stream);w.write(out)
    else:
        text={'.html':f'<html><script>secret()</script><p>{text}</p></html>','.htm':f'<p>{text}</p>','.json':json.dumps({'policy':text}),'.csv':f'policy\n"{text}"','.rtf':r'{\rtf1\ansi '+text+'}'}.get(ext,text)
        return text.encode()
    return out.getvalue()

@pytest.mark.parametrize('ext',['.txt','.md','.csv','.json','.html','.htm','.pdf','.docx','.pptx','.xlsx','.rtf'])
def test_upload_formats_and_search(client,ext):
    r=client.post('/api/documents',files={'file':('policy'+ext,payload(ext))});assert r.status_code==201,r.text
    result=client.post('/api/search',json={'query':'vacation','mode':'keyword'}).json()
    assert result['results'];assert 'vacation' in result['results'][0]['text']
    assert result['evidence'][0]['document']=='policy'+ext

def test_full_lifecycle(client):
    r=client.post('/api/documents',files={'file':('../../policy.txt',b'Car insurance protects your vehicle. Automobile coverage includes theft.')})
    identifier=r.json()['id'];assert client.get('/api/documents').json()[0]['name']=='policy.txt'
    for mode in ['keyword','semantic','hybrid']:
        result=client.post('/api/search',json={'query':'vehicle','mode':mode}).json()
        assert result['results'][0]['document_id']==identifier
    # The semantic synonym map can retrieve a term absent from the source.
    client.post('/api/documents',files={'file':('staff.txt',b'Employees may take paid holiday each year.')})
    exact=client.post('/api/search',json={'query':'vacation','mode':'keyword'}).json()
    related=client.post('/api/search',json={'query':'vacation','mode':'semantic'}).json()
    assert not exact['results'];assert related['results'][0]['name']=='staff.txt'
    assert client.delete('/api/documents/'+identifier).status_code==200
    assert client.delete('/api/documents/'+identifier).status_code==404
    assert not client.post('/api/search',json={'query':'vehicle','document_id':identifier}).json()['results']
    import app.main as main
    importlib.reload(main)
    assert len(TestClient(main.app).get('/api/documents').json())==1

def test_invalid_uploads(client):
    for name,data,status in [('x.exe',b'abc',415),('x.txt',b'',422),('x.pdf',b'not a PDF',422),('x.json',b'{broken',422),('x.txt',b'X'*(10*1024*1024+1),413)]:
        assert client.post('/api/documents',files={'file':(name,data)}).status_code==status
    assert client.post('/api/documents?chunk_size=60&overlap=60',files={'file':('x.txt',b'hello')}).status_code==422

@pytest.mark.parametrize('query',[{'query':''},{'query':' '},{'query':'x','top_k':0},{'query':'x','top_k':31},{'query':'x','alpha':2},{'query':'x','min_score':-1},{'query':'x','mode':'bad'}])
def test_query_validation(client,query): assert client.post('/api/search',json=query).status_code==422

def test_fusion_parameters_and_filters(client):
    for name,text in [('a.txt','Staff take holiday every summer.'),('b.txt','Vacation is a destination brand.')]:
        client.post('/api/documents',files={'file':(name,text.encode())})
    for alpha,mode in [(0,'keyword'),(1,'semantic')]:
        hybrid=client.post('/api/search',json={'query':'vacation','alpha':alpha}).json()['results']
        single=client.post('/api/search',json={'query':'vacation','mode':mode}).json()['results']
        assert [r['id'] for r in hybrid]==[r['id'] for r in single]
    docs=client.get('/api/documents').json()
    r=client.post('/api/search',json={'query':'vacation','document_id':docs[0]['id'],'top_k':1}).json()
    assert len(r['results'])==1 and r['results'][0]['document_id']==docs[0]['id']
    assert not client.post('/api/search',json={'query':'vacation','min_score':1}).json()['results']

def test_chunking_and_empty_search(client):
    text=' '.join(str(n) for n in range(300));parts=chunks(text,100,20)
    assert len(parts)==4 and parts[0].split()[-20:]==parts[1].split()[:20]
    assert not client.post('/api/search',json={'query':'anything'}).json()['results']
    assert client.get('/api/health').json()['status']=='ok'
    assert client.get('/').status_code==200

@pytest.mark.parametrize('error_type',['OperationalError','InvalidPassword','InsufficientPrivilege'])
def test_database_errors_are_actionable_and_do_not_leak_credentials(client,monkeypatch,caplog,error_type):
    import psycopg
    import app.main as main
    cls=getattr(psycopg,error_type,None) or getattr(psycopg.errors,error_type)
    secret='private-database-password-not-for-response'
    def unavailable(*args): raise cls(secret)
    monkeypatch.setattr(main,'connect',unavailable)
    r=client.post('/api/documents',files={'file':('policy.txt',b'Vacation policy for employees.')})
    assert r.status_code==503
    body=r.json()
    assert body['code']=='DATABASE_UNAVAILABLE' and len(body['reference'])==12
    assert 'database' in body['detail'].lower()
    assert secret not in r.text and secret not in caplog.text
    assert body['reference'] in caplog.text

def test_unexpected_error_has_safe_reference(client,monkeypatch,caplog):
    import app.main as main
    def broken(*args): raise RuntimeError('private internal connection details')
    monkeypatch.setattr(main,'connect',broken)
    with TestClient(main.app,raise_server_exceptions=False) as c:
        r=c.get('/api/health')
    assert r.status_code==500 and r.json()['code']=='INTERNAL_ERROR'
    assert 'private internal' not in r.text and 'private internal' not in caplog.text

def test_extracted_null_characters_are_removed(client,monkeypatch):
    import app.main as main
    monkeypatch.setattr(main,'extract',lambda *args:'Vacation\x00 policy for employees')
    r=client.post('/api/documents',files={'file':('policy.pdf',b'parser replaced by test')})
    assert r.status_code==201
    results=client.post('/api/search',json={'query':'Vacation','mode':'keyword'}).json()['results']
    assert results and '\x00' not in results[0]['text']

def test_postgres_does_not_prepare_statements():
    url=os.environ.get('TEST_DATABASE_URL')
    if not url: pytest.skip('Requires disposable PostgreSQL test database')
    from app.storage import PostgresConnection
    c=PostgresConnection(url)
    try:
        for _ in range(8): c.execute('SELECT 1').fetchone()
        assert c.execute('SELECT COUNT(*) AS count FROM pg_prepared_statements').fetchone()['count']==0
    finally: c.close()
