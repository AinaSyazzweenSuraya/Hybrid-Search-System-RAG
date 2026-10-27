"""Regression tests for identified relevance and formatting defects."""
from app.engine import retrieve,chunks
from app.parsers import extract
from docx import Document
import io
import pytest

@pytest.mark.parametrize('mode',['keyword','semantic','hybrid'])
def test_clinic_location_ranks_above_wrong_area_and_unrelated_document(mode):
    rows=[
        {'id':'wrong-area','name':'Panel clinics.txt','text':'Panel clinics in Shah Alam include Example Medical.'},
        {'id':'shopping','name':'Damansara shopping.txt','text':'I want to know the Damansara shopping area opening hours.'},
        {'id':'correct','name':'Medical directory.txt','text':'Panel clinics in Damansara include Example Damansara Clinic.'},
    ]
    found=retrieve(rows,'i want to know panel clinic in damansara area',mode=mode)
    assert found[0]['id']=='correct'
    assert not found[0]['missing_terms']
    assert next(r for r in found if r['id']=='wrong-area')['missing_terms']==['damansara']

@pytest.mark.parametrize('mode',['keyword','semantic','hybrid'])
def test_unanswerable_questions_do_not_match_only_filler_words(mode):
    rows=[{'id':'1','name':'Security.txt','text':'The policy is that employees must use multi-factor authentication.'}]
    assert not retrieve(rows,'What is the quantum flux capacitor warranty?',mode=mode)
    assert not retrieve(rows,'the and is',mode=mode)

@pytest.mark.parametrize('mode',['keyword','semantic','hybrid'])
def test_filename_is_searchable_without_rewriting_source_text(mode):
    text='Approved clinics include Example Medical Centre.'
    result=retrieve([{'id':'1','name':'Panel clinics managed by Micare.pdf','text':text}],'Micare',mode=mode)
    assert result and result[0]['text']==text

def test_missing_location_is_disclosed():
    r=retrieve([{'id':'1','name':'Panel clinic.txt','text':'Approved clinics are in Shah Alam.'}], 'panel clinic Damansara')
    assert r and 'damansara' in r[0]['missing_terms']

def test_chunking_keeps_source_line_breaks():
    text='Clinic name | Location\nExample Medical | Damansara\nExample Health | Shah Alam'
    assert chunks(text,60,10)==[text]

def test_docx_tables_remain_between_the_original_paragraphs():
    d=Document();d.add_paragraph('Before table');table=d.add_table(rows=1,cols=2)
    table.cell(0,0).text='Example Medical';table.cell(0,1).text='Damansara'
    d.add_paragraph('After table');buffer=io.BytesIO();d.save(buffer)
    text=extract('clinics.docx',buffer.getvalue())
    assert text.index('Before table')<text.index('Example Medical')<text.index('After table')
